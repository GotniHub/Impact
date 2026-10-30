"""Import sans exécuter les macros, formules ou instructions des fichiers."""
import csv
import io
import json
import re
import unicodedata
import zipfile
from html.parser import HTMLParser
from pathlib import Path

MAX_ROWS = 20000
MAX_BYTES = 25 * 1024 * 1024

def norm(value):
    return re.sub(r'[^a-z0-9]+', ' ', unicodedata.normalize('NFKD', str(value or '')).encode('ascii','ignore').decode().lower()).strip()

def clean(value):
    if value is None:
        return ''
    value = str(value).strip()
    if norm(value) in {'na','n a','n/a','none','null','nan','not available','not found','non disponible','inconnu','non renseigne','n d'} or value in {'-','—','/'}:
        return ''
    return value

ALIASES = {
    'name':['entreprise','company','company name','exposant','exhibitor','nom','societe','raison sociale'],
    'website':['website','site','site web','web','url site'],
    'category':['categorie','category','segment','secteur','sector'],
    'activity':['activite','activity','description'],
    'email':['email','e mail','mail','email entreprise'],
    'phone':['telephone','phone','tel','telephone entreprise'],
    'stand':['stand','booth','stand sigep','numero de stand','stand number','stand salon'],
    'hall':['hall','pavillon'], 'exhibitor_url':['sirha url','url exposant','exhibitor url','page exposant'],
    'city':['ville','city','implantation'], 'country':['pays','country'],
    'siren':['siren','siret'], 'company_linkedin':['linkedin entreprise','company linkedin'],
    'contact_name':['contact cle','nom contact','contact name','nom du directeur commercial ou chef de zone export'],
    'contact_role':['fonction','role','poste','contact role','directeur commercial ou chef de zone export'],
    'contact_email':['contact email','email contact'], 'contact_phone':['contact phone','telephone contact'],
    'contact_linkedin':['linkedin contact','contact linkedin'],
    'revenue_eur':['ca eur','ca','chiffre d affaires','revenue eur'],
    'margin_pct':['marge pct','marge operationnelle','margin pct'],
    'finance_year':['exercice','annee ca','finance year'],
    'finance_scope':['perimetre financier','finance scope'],
    'international':['international','export'], 'business_model':['business model','modele economique'],
    'sector_fit':['sector fit','adequation sectorielle'], 'employees':['effectif','employees'],
    'status':['statut','status','statut crm','type relation'],
}

def suggest_mapping(headers):
    result = {}
    for field, aliases in ALIASES.items():
        options = {norm(field), *map(norm, aliases)}
        hits = [h for h in headers if norm(re.sub(r' \(\d+\)$','',h)) in options]
        if hits:
            result[field] = hits[-1] if field == 'contact_linkedin' else hits[0]
    return result

def unique_headers(row):
    seen, result = {}, []
    for i, v in enumerate(row):
        v = clean(v) or f'Colonne {i+1}'
        seen[v] = seen.get(v,0)+1
        result.append(v if seen[v] == 1 else f'{v} ({seen[v]})')
    return result

def table(name, rows, header_row=0):
    if not rows:
        return None
    # Chercher une entête reconnue dans les dix premières lignes.
    header_row = max(range(min(10,len(rows))), key=lambda i:len(suggest_mapping([str(v or '') for v in rows[i]])))
    if all(len(row)==1 for row in rows) and not suggest_mapping([str(rows[header_row][0] or '')]).get('name'):
        rows=[['Entreprise']]+rows
        header_row=0
    headers = unique_headers(rows[header_row])
    records = []
    for number, row in enumerate(rows[header_row+1:],header_row+2):
        if any(clean(v) for v in row):
            records.append({'_row':number, **{h:clean(row[i]) if i<len(row) else '' for i,h in enumerate(headers)}})
    if len(records) > MAX_ROWS:
        raise ValueError(f'Maximum {MAX_ROWS} lignes par feuille.')
    mapping=suggest_mapping(headers)
    for field,aliases in ALIASES.items():
        options={norm(field),*map(norm,aliases)}
        hits=[h for h in headers if norm(re.sub(r' \(\d+\)$','',h)) in options]
        if len(hits)>1:
            # Certains anciens fichiers contiennent une colonne reconnue mais vide.
            mapping[field]=max(hits,key=lambda h:(sum(bool(r.get(h)) for r in records),headers.index(h) if field=='contact_linkedin' else -headers.index(h)))
    return {'name':name, 'headers':headers, 'rows':records, 'mapping':mapping}

def guard_zip(content):
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        if sum(x.file_size for x in z.infolist()) > 150*1024*1024:
            raise ValueError('Archive décompressée trop volumineuse.')

class HTMLTables(HTMLParser):
    def __init__(self):
        super().__init__(); self.tables=[]; self.rows=None; self.row=None; self.cell=None; self.text=[]; self.skip=0
    def handle_starttag(self, tag, attrs):
        if tag in {'script','style'}: self.skip += 1
        if tag == 'table': self.rows=[]
        elif tag=='tr' and self.rows is not None: self.row=[]
        elif tag in {'td','th'} and self.row is not None: self.cell=[]
    def handle_endtag(self,tag):
        if tag in {'script','style'}: self.skip=max(0,self.skip-1)
        if tag in {'td','th'} and self.cell is not None:
            self.row.append(' '.join(self.cell)); self.cell=None
        elif tag=='tr' and self.row is not None:
            self.rows.append(self.row); self.row=None
        elif tag=='table' and self.rows is not None:
            self.tables.append(self.rows); self.rows=None
    def handle_data(self,data):
        if not self.skip:
            self.text.append(data)
            if self.cell is not None: self.cell.append(data)

def parse_file(filename, content):
    if not content or len(content)>MAX_BYTES:
        raise ValueError('Le fichier doit contenir entre 1 octet et 25 Mo.')
    suffix=Path(filename).suffix.lower()
    tables, text, warnings = [], '', []
    if suffix in {'.xlsx','.xlsm'}:
        guard_zip(content)
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(content), read_only=False, data_only=True)
        for s in wb:
            if s.max_row>MAX_ROWS+20 or s.max_column>300:
                raise ValueError('Feuille trop grande (20 000 lignes ou 300 colonnes maximum).')
            tab=table(s.title,list(s.iter_rows(values_only=True)))
            if tab:
                for field in ['website','exhibitor_url','company_linkedin','contact_linkedin']:
                    header=tab['mapping'].get(field)
                    if header:
                        col=tab['headers'].index(header)+1
                        for row in tab['rows']:
                            cell=s.cell(row['_row'],col)
                            if cell.hyperlink and cell.hyperlink.target:
                                row[header]=cell.hyperlink.target
            tables.append(tab)
        wb.close()
        warnings.append('Les formules utilisent leur valeur enregistrée ; les macros ne sont jamais exécutées.')
    elif suffix=='.xls':
        import xlrd
        wb=xlrd.open_workbook(file_contents=content)
        tables=[table(s.name,[s.row_values(i) for i in range(s.nrows)]) for s in wb.sheets()]
    elif suffix in {'.csv','.tsv','.txt'}:
        for encoding in ['utf-8-sig','utf-16','cp1252']:
            try: text=content.decode(encoding); break
            except UnicodeError: continue
        else: raise ValueError('Encodage du texte non reconnu.')
        if suffix=='.tsv': delimiter='\t'
        else:
            try: delimiter=csv.Sniffer().sniff(text[:10000],delimiters=',;\t').delimiter
            except csv.Error: delimiter=None
        if delimiter:
            tables=[table('Données',list(csv.reader(io.StringIO(text),delimiter=delimiter)))]
        elif suffix=='.csv':
            tables=[table('Données',list(csv.reader(io.StringIO(text))))]
        elif suffix=='.txt':
            tables=[table('Noms', [['Entreprise']]+[[l.strip()] for l in text.splitlines() if l.strip()])]
    elif suffix=='.json':
        data=json.loads(content.decode('utf-8-sig'))
        if isinstance(data,dict):
            data=next((data[k] for k in ['exhibitors','exposants','companies','data','results'] if isinstance(data.get(k),list)),data)
        if not isinstance(data,list): raise ValueError('JSON attendu : une liste ou une clé exhibitors/exposants/data/results.')
        if len(data)>MAX_ROWS: raise ValueError('Trop de lignes JSON.')
        if all(isinstance(x,str) for x in data): tables=[table('Données',[['Entreprise']]+[[x] for x in data])]
        elif all(isinstance(x,dict) for x in data):
            headers=list(dict.fromkeys(k for x in data for k in x))
            tables=[table('Données',[headers]+[[json.dumps(x.get(h),ensure_ascii=False) if isinstance(x.get(h),(dict,list)) else x.get(h,'') for h in headers] for x in data])]
        else: raise ValueError('Chaque élément JSON doit être un nom ou un objet.')
    elif suffix=='.docx':
        guard_zip(content)
        from docx import Document
        doc=Document(io.BytesIO(content))
        tables=[table(f'Tableau {i+1}',[[c.text for c in row.cells] for row in t.rows]) for i,t in enumerate(doc.tables)]
        text='\n'.join(p.text for p in doc.paragraphs)
    elif suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(content))
        if len(reader.pages)>300: raise ValueError('Maximum 300 pages PDF.')
        text='\n'.join(page.extract_text() or '' for page in reader.pages)
        warnings.append('PDF : vérifier les associations entre entreprise, hall et stand après extraction IA.')
    elif suffix in {'.html','.htm'}:
        parser=HTMLTables(); parser.feed(content.decode('utf-8',errors='replace'))
        tables=[table(f'Tableau {i+1}',t) for i,t in enumerate(parser.tables)]
        text=' '.join(parser.text)
    else:
        raise ValueError('Format non pris en charge. Utiliser XLSX, XLS, CSV, TSV, JSON, DOCX, PDF texte, TXT ou HTML.')
    tables=[t for t in tables if t and t['rows']]
    if not tables and not text.strip(): raise ValueError('Aucune donnée lisible. Un document scanné doit être converti par OCR avant import.')
    if len(text)>200000:
        raise ValueError('Texte trop long. Diviser le document ou importer une table CSV/XLSX.')
    return {'filename':Path(filename).name,'tables':tables,'text':text,'warnings':warnings}

def normalize_records(tab, mapping):
    if not mapping.get('name'):
        raise ValueError('Associer la colonne Entreprise avant de continuer.')
    if any(h not in tab['headers'] for h in mapping.values() if h):
        raise ValueError('Une colonne associée est absente de cette feuille.')
    result, warnings, previous = [], [], None
    for row in tab['rows']:
        record={k:clean(row.get(h)) for k,h in mapping.items() if h}
        record['_row']=row['_row']
        if not record.get('name'):
            if previous and any(record.get(k) for k in ['contact_name','contact_email','contact_linkedin']):
                # SIGEP : lignes de contacts sous l'entreprise, jamais sous une autre société.
                inherited={k:v for k,v in previous.items() if not k.startswith('contact_') and k not in {'_row'}}
                record={**inherited,**{k:v for k,v in record.items() if v or k=='_row'}}
            else:
                warnings.append(f'Ligne {row["_row"]} ignorée : entreprise absente.'); continue
        else: previous=record.copy()
        result.append(record)
    if not result: raise ValueError('Aucune entreprise importable avec cette association de colonnes.')
    return result,warnings
