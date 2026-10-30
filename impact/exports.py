import csv
import hashlib
import io
import json
import re
import zipfile
from html import escape
from .config import FIELDS, TRIGGERS
from .domain import value, contact_rows

EXPORT_MODES={
    'Présélection automatique':{'validated','proposed','shortlisted'},
    'Critères documentés':{'validated','proposed'},
    'Vérifiées manuellement':{'validated'},
}
TIER_LABELS={1:'Très forte priorité',2:'Forte priorité',3:'Priorité secondaire'}


def tier_title(tier):
    return f'Tier {tier} {TIER_LABELS[tier]}'


def sorted_memos(companies,confirmed=False):
    key='validated_assessment' if confirmed else 'assessment'
    return sorted(companies,key=lambda c:(c[key]['tier'],-c[key]['score'],c['name'].casefold()))


def selected_for_export(companies,mode):
    if mode not in EXPORT_MODES: raise ValueError('Choix d’export inconnu.')
    return sorted_memos([c for c in companies if c['selection'] in EXPORT_MODES[mode]],mode=='Vérifiées manuellement')


def export_signature(campaign,companies,crm,mode):
    """Empêche de servir un paquet préparé avant une modification des données."""
    material={'campaign':campaign,'companies':companies,'crm':crm,'mode':mode}
    return hashlib.sha256(json.dumps(material,sort_keys=True,ensure_ascii=False).encode('utf-8')).hexdigest()


def prepare_package(campaign,companies,crm,mode):
    confirmed=mode=='Vérifiées manuellement'
    selected=selected_for_export(companies,mode)
    package={'campaign_id':campaign['id'],'confirmed':confirmed,'mode':mode,
             'signature':export_signature(campaign,companies,crm,mode),
             'xlsx':workbook_bytes(campaign,companies,crm),'csv':csv_bytes(contact_rows(selected,confirmed)),
             'docx':None,'pdf':None,'zip':None,'errors':[]}
    if not selected: return package
    # Word/PDF facultatifs : une dépendance manquante ne doit pas bloquer Excel et CSV.
    for name,label,builder in [('docx','Word',document_bytes),('pdf','PDF',pdf_bytes)]:
        try: package[name]=builder(campaign,selected,confirmed)
        except ImportError as exc:
            module=getattr(exc,'name',None) or 'dépendance d’export'
            distribution={'docx':'python-docx','reportlab':'reportlab'}.get(module,module)
            package['errors'].append(f'Export {label} indisponible dans le Python utilisé : {distribution} manque. '
                                     'Relancez l’application avec Lancer_IMPACT.cmd, qui utilise l’environnement du projet. '
                                     'Les autres formats disponibles restent téléchargeables.')
    if package['docx'] is not None and package['pdf'] is not None:
        bundle_companies=[c for c in companies if c['selection']!='shortlisted'] if mode=='Critères documentés' else companies
        package['zip']=bundle_bytes(campaign,bundle_companies,crm,confirmed)
    return package

def safe_text(v):
    text=str(v if v is not None else '')
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) else text

def safe_filename(name):
    return re.sub(r'[^\w\- ]','',name,flags=re.UNICODE).strip()[:80] or 'entreprise'

def csv_bytes(rows):
    output=io.StringIO(newline='')
    headers=list(rows[0]) if rows else ['Entreprise','Qualification']
    writer=csv.DictWriter(output,fieldnames=headers,delimiter=';'); writer.writeheader()
    writer.writerows({k:safe_text(v) for k,v in row.items()} for row in rows)
    return output.getvalue().encode('utf-8-sig')

def summary_rows(companies,confirmed=False):
    rows=[]
    for c in companies:
        a=c['validated_assessment'] if confirmed else c['assessment']
        rows.append({'Entreprise':c['name'],'Tier':a['tier'],'Score':a['score'],'Statut':c['selection'],
                     'Stand':value(c,'stand',confirmed) or 'Non renseigné','Hall':value(c,'hall',confirmed) or '',
                     'Activité':value(c,'activity',confirmed) or value(c,'category',confirmed) or '',
                     'Site web':value(c,'website',confirmed) or '',
                     'CA EUR':value(c,'revenue_eur',confirmed),'Marge opérationnelle %':value(c,'margin_pct',confirmed),
                     'Exercice':value(c,'finance_year',confirmed),'Périmètre':value(c,'finance_scope',confirmed) or '',
                     'Points à documenter':' ; '.join(a['blockers']),
                     'Vérification humaine':'Effectuée' if c['selection']=='validated' else 'À faire'})
    return rows

def aggregate_rows(crm,companies):
    rows=[]; seen=set()
    for company in crm:
        for contact in company.get('contacts') or [{}]:
            row={'Entreprise':company['name'],'Statut CRM':company['status'],'Segment':company.get('category',''),
                 'Site web':company.get('website',''),'SIREN':company.get('siren',''),
                 'Email entreprise':company.get('email',''),'Téléphone entreprise':company.get('phone',''),
                 'Nom du contact':contact.get('name',''),'Fonction':contact.get('role',''),
                 'Email contact':contact.get('email',''),'Téléphone contact':contact.get('phone',''),
                 'LinkedIn contact':contact.get('linkedin',''),'Source':company.get('source','')}
            signature=(company['name'].lower(),contact.get('name','').lower(),contact.get('email','').lower())
            if signature not in seen: rows.append(row); seen.add(signature)
    for c in companies:
        if c.get('selection')!='validated': continue
        for contact in [x for x in c.get('contacts',[]) if x.get('status')=='validated'] or [{}]:
            signature=(c['name'].lower(),contact.get('name','').lower(),contact.get('email','').lower())
            if signature in seen: continue
            seen.add(signature)
            rows.append({'Entreprise':c['name'],'Statut CRM':'prospect à intégrer','Segment':value(c,'category',True) or '',
                         'Site web':value(c,'website',True) or '', 'SIREN':value(c,'siren',True) or '',
                         'Email entreprise':value(c,'email',True) or '','Téléphone entreprise':value(c,'phone',True) or '',
                         'Nom du contact':contact.get('name',''),'Fonction':contact.get('role',''),
                         'Email contact':contact.get('email',''),'Téléphone contact':contact.get('phone',''),
                         'LinkedIn contact':contact.get('linkedin',''),'Source':contact.get('source','')})
    return rows

def workbook_bytes(campaign,companies,crm):
    # Générateur intégré à l'application ; les données restent typées et les chaînes ne deviennent pas des formules.
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook(); wb.remove(wb.active)
    accepted=[c for c in companies if c['selection']=='validated']
    proposed=[c for c in companies if c['selection']=='proposed']
    shortlisted=[c for c in companies if c['selection']=='shortlisted']
    backup=[c for c in companies if c['selection']=='backup']
    source_rows=[]
    for c in companies:
        for field,fact in c.get('facts',{}).items():
            source_rows.append({'Entreprise':c['name'],'Champ':FIELDS.get(field,field),'Valeur':fact['value'],
                                'Statut':fact.get('status',''),'Source':fact.get('source',''),'Extrait':fact.get('quote',''),
                                'Exercice':fact.get('year'),'Périmètre':fact.get('scope',''),'Vérifié le':fact.get('checked_at','')})
        for t in c.get('triggers',[]):
            source_rows.append({'Entreprise':c['name'],'Champ':t['code'],'Valeur':t.get('observation',''),'Statut':t.get('status',''),
                                'Source':t.get('source',''),'Extrait':t.get('quote',''),'Exercice':None,'Périmètre':'','Vérifié le':t.get('checked_at','')})
    excluded=[{'Entreprise':c['name'],'Décision':c['assessment']['exclusion']['state'],'Raison':c['assessment']['exclusion']['reason']} for c in companies if c['assessment']['exclusion']['state']!='clear']
    settings=[{'Paramètre':'Salon','Valeur':campaign['name']},{'Paramètre':'Édition','Valeur':campaign['edition']},
              {'Paramètre':'Base CRM','Valeur':str(len(crm))+' entreprises'},
              {'Paramètre':'Cibles validées','Valeur':f'{len(accepted)} / {campaign["criteria"]["target_count"]}'}]
    settings += [{'Paramètre':k,'Valeur':json.dumps(v,ensure_ascii=False)} for k,v in campaign['criteria'].items()]
    dismissed=[{'Entreprise':c['name'],'Raison':c['selection_decision'].get('reason',''),
                'Décision le':c['selection_decision'].get('updated','')} for c in companies
               if c.get('selection_decision',{}).get('dismissed')]
    tables=[('Cibles validées',summary_rows(accepted,True)),('Propositions IA',summary_rows(proposed)),
            ('Présélection à documenter',summary_rows(shortlisted)),('Réserve',summary_rows(backup)),
            ('Contacts',contact_rows(accepted,True)),('Base agrégée',aggregate_rows(crm,companies)),('Contacts à vérifier',contact_rows(proposed+shortlisted)),('Sources',source_rows),('Exclusions',excluded),('Écartées de la sélection',dismissed),('Paramètres',settings)]
    for title,rows in tables:
        sheet=wb.create_sheet(title)
        headers=list(rows[0]) if rows else (['Entreprise','Statut'] if title!='Paramètres' else ['Paramètre','Valeur'])
        sheet.append(headers)
        for row in rows:
            sheet.append([safe_text(row.get(h)) if isinstance(row.get(h),str) else row.get(h) for h in headers])
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        sheet.sheet_view.showGridLines=False
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='243C37'); cell.font=Font(name='Calibri',bold=True,color='FFFFFF'); cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[1].height=32
        for i,h in enumerate(headers,1):
            width=55 if h in {'Activité','Points à vérifier','Extrait','Source','Valeur'} else 35 if 'web' in h or 'LinkedIn' in h else 28 if h in {'Entreprise','Champ','Fonction','Nom du contact'} else 20
            sheet.column_dimensions[get_column_letter(i)].width=width
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height=54 if title=='Sources' else 42
            for cell in row:
                cell.font=Font(name='Calibri',size=11,color='253A34'); cell.alignment=Alignment(wrap_text=True,vertical='top')
                if cell.row%2==0: cell.fill=PatternFill('solid',fgColor='F1F5F3')
                if headers[cell.column-1]=='CA EUR': cell.number_format='#,##0.00'
                if headers[cell.column-1] in {'Marge opérationnelle %','Score'}: cell.number_format='0.0'
    output=io.BytesIO(); wb.save(output); return output.getvalue()

def memo_sections(company,confirmed=False):
    v=lambda f:value(company,f,confirmed)
    facts=company['facts']; memo=company.get('memo',{})
    pitch_lines=[memo.get('pitch') or 'Explorer ensemble les arbitrages de gamme et la contribution des références par circuit']
    if memo.get('generation_mode')=='local_standard':
        pitch_lines.insert(0,'Support local standard à personnaliser ; les besoins et les informations sourcées restent à vérifier.')
    metrics=[]
    for field in ['revenue_eur','margin_pct','employees']:
        n=v(field)
        if n is None: metrics.append(f'{FIELDS[field]} : non renseigné'); continue
        fact=facts[field]
        unit=' EUR' if field=='revenue_eur' else ' %' if field=='margin_pct' else ''
        text=(f'{FIELDS[field]} : {n:,.1f}{unit}' if field=='margin_pct' else f'{FIELDS[field]} : {n:,.0f}{unit}').replace(',',' ')
        if field in {'revenue_eur','margin_pct'}: text+=f' · exercice {fact.get("year") or v("finance_year") or "inconnu"} · {fact.get("scope") or v("finance_scope") or "périmètre inconnu"}'
        metrics.append(text)
    from .signals import partition
    triggers,pending_triggers=partition(company,confirmed)
    events=[e for e in company.get('events',[]) if not confirmed or e.get('status')=='validated']
    contacts=[c for c in company.get('contacts',[]) if not confirmed or c.get('status')=='validated']
    sources=[]
    for field,fact in facts.items():
        if not confirmed or fact.get('status')=='validated':
            sources.append(f'{FIELDS.get(field,field)} : {fact.get("source","")}')
    for t in triggers: sources.append(f'{t["code"]} : {t.get("source","")}')
    for e in events: sources.append(f'Évolution : {e.get("source","")}')
    for c in contacts: sources.append(f'Contact {c.get("name","")} : {c.get("source","")}')
    qualification=[]
    if company.get('selection')=='shortlisted':
        qualification=[('Qualification',['Présélection automatique à documenter. Les critères obligatoires ne sont pas tous démontrés.']+
                        company.get('assessment',{}).get('blockers',[]))]
    pending_sections=[]
    if pending_triggers:
        pending_sections=[('Signaux proposés à vérifier',[f'{t.get("code")} : {reason} Ne compte pas dans le score.' for t,reason in pending_triggers])]
    return qualification+[
        ('Qui est cette entreprise',[v('activity') or v('category') or 'Activité non renseignée',
                                    f'Modèle : {v("business_model") or "non renseigné"}. Implantation : {v("city") or "non renseignée"}.']),
        ('Évolutions documentées',[f'{e["year"]} — {e["event"]}' for e in events] or ['Aucune évolution documentée']),
        ('Chiffres clés',metrics),
        ('Signaux SENSO',[f'{t["code"]} — {TRIGGERS[t["code"]]["name"]}. {t.get("observation","")}' for t in triggers] or ['Aucun signal documenté']),
    ]+pending_sections+[
        ('Hypothèses de besoins à confirmer',memo.get('pain_points') or ['À établir pendant l’échange ; aucune douleur interne supposée comme un fait']),
        ('Proposition de pitch',pitch_lines),
        ('Questions pour le rendez vous',memo.get('questions') or ['Comment décidez vous quelles références lancer, maintenir ou arrêter ?','Quels circuits ou gammes nécessitent le plus d’arbitrages ?']),
        ('Contacts',[f'{c.get("name") or "Contact générique"} · {c.get("role","")} · {c.get("email","")} · {c.get("phone","")} · {c.get("linkedin","")}' for c in contacts] or [f'Email entreprise : {v("email") or "non renseigné"} · Téléphone : {v("phone") or "non renseigné"}']),
        ('Phrase clé à retenir',[memo.get('key_message') or 'Confirmer les besoins avant de proposer un arbitrage de portefeuille']),
        ('Sources',list(dict.fromkeys(sources))),
    ]

def document_bytes(campaign,companies,confirmed=False):
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    doc=Document()
    section=doc.sections[0]; section.top_margin=section.bottom_margin=Cm(1.8); section.left_margin=section.right_margin=Cm(2)
    for style in ['Normal','Title','Heading 1','Heading 2']:
        doc.styles[style].font.name='Calibri'; doc.styles[style].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(4)
    doc.styles['Title'].font.size=Pt(22)
    doc.styles['Heading 1'].font.size=Pt(13)
    doc.core_properties.title=f'Fiches mémos {campaign["name"]} {campaign["edition"]}'
    previous_tier=None
    for i,c in enumerate(sorted_memos(companies,confirmed)):
        if i: doc.add_page_break()
        a=c['validated_assessment'] if confirmed else c['assessment']
        if a['tier']!=previous_tier:
            doc.add_heading(tier_title(a['tier']),1)
            if i==0: doc.add_paragraph('Priorité calculée sur les informations disponibles. Les données manquantes et les signaux non confirmés restent à vérifier.')
            previous_tier=a['tier']
        doc.add_paragraph(c['name'],style='Title')
        doc.add_paragraph(f'{campaign["name"]} · {campaign["edition"]}')
        a=c['validated_assessment'] if confirmed else c['assessment']
        p=doc.add_paragraph(); r=p.add_run(f'STAND {value(c,"stand",confirmed) or "NON RENSEIGNÉ"}    HALL {value(c,"hall",confirmed) or "NON RENSEIGNÉ"}'); r.bold=True; r.font.size=Pt(15)
        doc.add_paragraph(f'Tier {a["tier"]} · Score {a["score"]}/100 · '+('Cible validée' if confirmed else 'Présélection à documenter' if c.get('selection')=='shortlisted' else 'Critères documentés, à vérifier'))
        for title,lines in memo_sections(c,confirmed):
            doc.add_heading(title,2)
            for line in lines:
                if title=='Sources' and 'https://' in line or title=='Sources' and 'http://' in line:
                    match=re.search(r'https?://\S+',line)
                    p=doc.add_paragraph(line[:match.start()])
                    hyperlink=OxmlElement('w:hyperlink')
                    hyperlink.set(qn('r:id'),p.part.relate_to(match.group(),RT.HYPERLINK,is_external=True))
                    run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
                    size=OxmlElement('w:sz'); size.set(qn('w:val'),'16'); props.append(size)
                    run.append(props); text=OxmlElement('w:t'); text.text=match.group(); run.append(text); hyperlink.append(run); p._p.append(hyperlink)
                else:
                    p=doc.add_paragraph(line)
                if title=='Sources':
                    for run in p.runs: run.font.size=Pt(8)
    if not companies:
        doc.add_paragraph('Fiches mémos',style='Title'); doc.add_paragraph('Aucune cible sélectionnée.')
    output=io.BytesIO(); doc.save(output); return output.getvalue()

def pdf_bytes(campaign,companies,confirmed=False):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.enums import TA_LEFT
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='MemoTitle',fontName='Helvetica-Bold',fontSize=21,leading=25,spaceAfter=10))
    styles.add(ParagraphStyle(name='MemoTier',fontName='Helvetica-Bold',fontSize=14,leading=18,spaceAfter=9,keepWithNext=True))
    styles.add(ParagraphStyle(name='Stand',fontName='Helvetica-Bold',fontSize=14,leading=18,spaceAfter=8))
    styles.add(ParagraphStyle(name='MemoHeading',fontName='Helvetica-Bold',fontSize=11,leading=14,spaceBefore=9,spaceAfter=4,keepWithNext=True))
    styles.add(ParagraphStyle(name='MemoBody',fontName='Helvetica',fontSize=9,leading=12,spaceAfter=4))
    styles.add(ParagraphStyle(name='SourceText',fontName='Helvetica',fontSize=7,leading=9,spaceAfter=3,wordWrap='CJK'))
    story=[]
    def p(text,style):
        # XML/HTML provenant des données est toujours échappé.
        rendered=escape(str(text)).replace('\n','<br/>')
        if style=='SourceText':
            rendered=re.sub(r'https?://[^\s<]+',lambda m:f'<link href="{m.group()}">{m.group()}</link>',rendered)
        return Paragraph(rendered,styles[style])
    previous_tier=None
    for i,c in enumerate(sorted_memos(companies,confirmed)):
        if i: story.append(PageBreak())
        a=c['validated_assessment'] if confirmed else c['assessment']
        if a['tier']!=previous_tier:
            story.append(p(tier_title(a['tier']),'MemoTier'))
            if i==0: story.append(p('Priorité calculée sur les informations disponibles. Les données manquantes et les signaux non confirmés restent à vérifier.','MemoBody'))
            previous_tier=a['tier']
        story.append(p(c['name'],'MemoTitle')); story.append(p(f'{campaign["name"]} · {campaign["edition"]}','MemoBody'))
        story.append(p(f'STAND {value(c,"stand",confirmed) or "NON RENSEIGNÉ"}    HALL {value(c,"hall",confirmed) or "NON RENSEIGNÉ"}','Stand'))
        a=c['validated_assessment'] if confirmed else c['assessment']
        story.append(p(f'Tier {a["tier"]} · Score {a["score"]}/100 · '+('Cible validée' if confirmed else 'Présélection à documenter' if c.get('selection')=='shortlisted' else 'Critères documentés, à vérifier'),'MemoBody'))
        for title,lines in memo_sections(c,confirmed):
            story.append(p(title,'MemoHeading'))
            for line in lines: story.append(p(line,'SourceText' if title=='Sources' else 'MemoBody'))
    if not story: story=[p('Aucune cible sélectionnée','MemoTitle')]
    output=io.BytesIO()
    SimpleDocTemplate(output,pagesize=A4,rightMargin=45,leftMargin=45,topMargin=40,bottomMargin=40,title='Fiches mémos IMPACT').build(story)
    return output.getvalue()

def bundle_bytes(campaign,companies,crm,confirmed=False):
    selected=[c for c in companies if c['selection']=='validated' or (not confirmed and c['selection'] in {'proposed','shortlisted'})]
    selected=sorted_memos(selected,confirmed)
    if not selected: raise ValueError('Aucune cible à exporter dans ce mode.')
    output=io.BytesIO()
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('Prospection_IMPACT.xlsx',workbook_bytes(campaign,companies,crm))
        z.writestr('Contacts.csv',csv_bytes(contact_rows(selected,confirmed)))
        z.writestr('Base_agregee.csv',csv_bytes(aggregate_rows(crm,companies)))
        z.writestr('Fiches_Memos.docx',document_bytes(campaign,selected,confirmed))
        z.writestr('Fiches_Memos.pdf',pdf_bytes(campaign,selected,confirmed))
        for i,c in enumerate(selected,1):
            base=f'fiches/{i:02d}_{safe_filename(c["name"])}'
            z.writestr(base+'.docx',document_bytes(campaign,[c],confirmed))
            z.writestr(base+'.pdf',pdf_bytes(campaign,[c],confirmed))
        z.writestr('Resultats_et_sources.json',json.dumps({'campaign':campaign,'companies':selected},ensure_ascii=False,indent=2))
    return output.getvalue()
