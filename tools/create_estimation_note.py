"""Note de décision pour le chef de projet, avec hypothèses de coût explicites."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs' / 'Comparatif_IMPACT_OpenAI_Ollama_4000_exposants.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.75)
sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.75)
sec.footer_distance = Inches(.3)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Calibri')
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'), 'Calibri')
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = 1
doc.styles['Title'].font.size = Pt(22)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(6)
doc.styles['Subtitle'].font.size = Pt(13)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].font.bold = True
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)

def paragraph(text, bold_lead=None, style=None, small=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    if small:
        for r in p.runs:
            r.font.size = Pt(9.5)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(5)
    return p

def heading(text):
    return doc.add_heading(text, level=1)

def table(headers, rows, widths, numeric_cols=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Inches(width)
    tblpr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + edge)
        for k, v in {'val':'single', 'sz':'4', 'color':'D9D9D9'}.items():
            b.set(qn('w:' + k), v)
        borders.append(b)
    tblpr.append(borders)
    for values in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text
    for i, row in enumerate(t.rows):
        props = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        props.append(no_split)
        if i == 0:
            props.append(OxmlElement('w:tblHeader'))
            for cell, text in zip(row.cells, headers):
                cell.text = text
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '243E54' if i == 0 else ('F1F5F8' if i % 2 == 0 else 'FFFFFF'))
            tcpr.append(shading)
            margins = OxmlElement('w:tcMar')
            for edge, val in [('top', 65), ('bottom', 65), ('left', 110), ('right', 110)]:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:w'), str(val))
                e.set(qn('w:type'), 'dxa')
                margins.append(e)
            tcpr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.03
                if j in numeric_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = 'Calibri'
                    r.font.size = Pt(10.5)
                    r.font.bold = i == 0
                    r.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else RGBColor(0, 0, 0)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = 1
    after.add_run().font.size = Pt(3)
    return t

def link(p, label, url):
    rel = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rel)
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '244A70')
    pr.append(color)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '19')
    pr.append(sz)
    r.append(pr)
    text = OxmlElement('w:t')
    text.text = label
    r.append(text)
    h.append(r)
    p._p.append(h)

# Synthèse courte : comparaison et budget lisibles sans explication technique.
doc.add_paragraph('Budget du projet IMPACT', 'Title')
doc.add_paragraph('OpenAI et Ollama pour 4 000 exposants', 'Subtitle')
paragraph('Chef de projet | 9 octobre 2026 | Dollars américains ($)', small=True)
paragraph('À partir du fichier fourni : analyser les entreprises, exclure les clients et prospects connus, puis produire 20 cibles prioritaires avec contacts, stands et mémos.')

heading('Comparaison des solutions')
table(['Critère', 'OpenAI avec GPT', 'Ollama'], [
    ['Utilisation', 'IA en ligne.', 'IA sur le PC ou en ligne.'],
    ['Démarrage', 'Connexion déjà développée.', 'Adaptation nécessaire.'],
    ['Recherche web', 'Prévue dans le traitement.', 'Outil de recherche à relier.'],
    ['Coût', 'Paiement selon l’usage.', 'IA gratuite sur le PC ; frais web et cloud.'],
    ['Qualité et délai', 'À mesurer sur un essai.', 'Dépend du modèle et du PC en local.'],
], [1.35, 2.75, 2.90])

heading('Budget indicatif pour 4 000 exposants')
paragraph('Exemple : 4 000 entreprises analysées, avec 5 recherches web chacune, soit 20 000 recherches.', small=True)
table(['Solution', 'Analyse IA', 'Recherche web', 'Total estimé'], [
    ['OpenAI GPT 5.5', '1 000 $', '200 $', '1 200 $'],
    ['Ollama sur le PC + Serper', '0 $', '50 $', '50 $'],
    ['Ollama Cloud gpt oss 120b + Serper', '24 $', '50 $', '74 $'],
], [3.10, 1.10, 1.40, 1.40], numeric_cols=(1, 2, 3))
paragraph('OpenAI : environ 680 à 2 400 $ selon la profondeur, sans plafond garanti. Serper : pack de 50 000 recherches à 50 $. Coûts estimés, encore non mesurés.', small=True)

heading('Points à retenir avant de décider')
table(['Point', 'À retenir'], [
    ['Frais exclus', 'Taxes, développement, matériel, électricité et vérification humaine.'],
    ['Qualité', 'Vérifier les sources des chiffres, contacts et stands. Qualité identique entre modèles à confirmer.'],
    ['Recommandation', 'Tester les deux solutions sur 20 entreprises : coût, durée et exactitude avant les 4 000.'],
], [1.35, 5.65])

paragraph('Base de calcul : 20 000 tokens lus et 5 000 produits par entreprise, tous appels inclus. Calcul local gratuit ne signifie pas processus complet gratuit.', small=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1
r = p.add_run('Tarifs officiels consultés le 9 octobre 2026 : ')
r.font.size = Pt(9.5)
sources = [
    ('OpenAI modèle', 'https://developers.openai.com/api/docs/models/gpt-5.5'),
    ('OpenAI web', 'https://developers.openai.com/api/docs/pricing'),
    ('Ollama', 'https://ollama.com/pricing'),
    ('Serper', 'https://serper.dev/'),
]
for i, (label, url) in enumerate(sources):
    if i:
        p.add_run(' | ').font.size = Pt(9.5)
    link(p, label, url)

# Retirer aussi les bordures héritées de la police de styles fournie au runtime.
for style in doc.styles:
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
for p in doc.paragraphs:
    for border in list(p._p.xpath('.//w:pBdr')):
        border.getparent().remove(border)
doc.core_properties.title = 'Estimation du projet IMPACT OpenAI et Ollama pour 4 000 exposants'
doc.core_properties.subject = 'Comparaison et budget indicatif pour le chef de projet'
doc.core_properties.author = 'IMPACT'
doc.core_properties.keywords = 'IMPACT, OpenAI, Ollama, estimation, exposants'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
