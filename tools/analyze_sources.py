import json
from pathlib import Path
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FILES = [
    Path(r'C:/Users/IlyassGotni/Desktop/Sirha scrapper/converter marseille/Sirha_Marseille_Exhibitors_List.xlsx'),
    Path(r'C:/Users/IlyassGotni/Documents/Salon SIRHA MARSEILLE 20 cibles prio Fiche Mémos.docx'),
    Path(r'C:/Users/IlyassGotni/Documents/Triggers.docx'),
    Path(r'C:/Users/IlyassGotni/Downloads/SIGEP_Prospection_20_Cibles_Contact.xlsx'),
    Path(r'C:/Users/IlyassGotni/Documents/PROJECT IMPACT/PROCESSUS PROJECT IMPACT.docx'),
]
out = ROOT / 'analysis'
out.mkdir(exist_ok=True)
result = []
for path in FILES:
    if not path.exists():
        result.append({'file': str(path), 'missing': True})
        continue
    if path.suffix == '.docx':
        document = Document(path)
        lines = []
        for element in document.element.body:
            if element.tag.endswith('}p'):
                lines.append(''.join(element.itertext()))
            elif element.tag.endswith('}tbl'):
                for row in element.findall('./{*}tr'):
                    lines.append(' | '.join(''.join(cell.itertext()) for cell in row.findall('./{*}tc')))
        # python-docx paragraphs preserve the displayed text without OOXML duplicate itertext.
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        lines = []
        for element in document.element.body:
            if element.tag.endswith('}p'):
                lines.append(Paragraph(element, document).text)
            elif element.tag.endswith('}tbl'):
                lines += [' | '.join(cell.text for cell in row.cells) for row in Table(element, document).rows]
        text = '\n'.join(lines)
        (out / (path.stem + '.txt')).write_text(text, encoding='utf-8')
        result.append({'file':str(path),'text':text})
    else:
        wb = load_workbook(path, data_only=True, read_only=True)
        sheets = []
        for sheet in wb:
            rows = [[str(v) if v is not None else '' for v in row] for row in sheet.iter_rows(values_only=True)]
            sheets.append({'name':sheet.title,'rows':rows,'count':len(rows),'columns':sheet.max_column})
        result.append({'file':str(path),'sheets':sheets})
(out / 'sources.json').write_text(json.dumps(result, ensure_ascii=False, indent=2),encoding='utf-8')
for entry in result:
    if 'text' in entry:
        print(entry['file'], '\n', entry['text'], '\n')
    elif 'sheets' in entry:
        print(json.dumps({**entry, 'sheets':[{**s,'rows':s['rows'][:5]} for s in entry['sheets']]},ensure_ascii=False))
    else:
        print(entry)
