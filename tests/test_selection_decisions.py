import copy
import io
import json
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from impact import db,exports,services
from impact.domain import rank


ACTIVE={'validated','proposed','shortlisted'}


class SelectionDecisionTests(unittest.TestCase):
    def setUp(self):
        temp=tempfile.TemporaryDirectory(prefix='impact_selection_')
        self.addCleanup(temp.cleanup)
        data=patch.object(db,'DATA',Path(temp.name)); data.start(); self.addCleanup(data.stop)
        db.init()
        self.campaign=services.create_campaign('Salon décision','2027',[
            {'name':f'Aliment {i:02d}','category':'Produits alimentaires','sector_fit':'true',
             'business_model':'industriel','website':f'https://aliment{i}.example'} for i in range(10)],'test')
        self.campaign['criteria'].update(target_count=3,backup_count=4,progressive_selection=True)
        self.campaign['crm_review']={'empty_confirmed':True}; db.put('campaigns',self.campaign)
        for i,c in enumerate(db.all_rows('companies',self.campaign['id'])):
            c['research_status']='completed'
            for f in c['facts'].values(): f['status']='sourced'
            if i<2:
                for field,value in {'identity':True,'revenue_eur':12_000_000,
                                    'finance_year':datetime.now().year-1,'finance_scope':'entite'}.items():
                    c['facts'][field]={'value':value,'status':'sourced','source':'https://example.com'}
            db.put('companies',c)

    def test_replacement_is_persistent_reversible_and_preserves_research_and_crm(self):
        campaign,rows=services.campaign_results(self.campaign['id'])
        top=[c for c in rows if c['selection'] in ACTIVE]
        self.assertEqual([c['selection'] for c in top],['proposed','proposed','shortlisted'])
        ids=[top[0]['id'],top[2]['id']]
        before=db.all_rows('companies'); jobs=db.all_rows('jobs'); crm=db.all_rows('crm')
        with patch('impact.ai.research') as ai,patch('impact.services.start_job') as launch:
            change=services.set_selection_dismissed(campaign['id'],ids,reason='Non adapté au projet')
            updated,after=services.campaign_results(campaign['id'])
            selected=[c for c in after if c['selection'] in ACTIVE]
            self.assertEqual(len(selected),3); self.assertEqual(len(change['added']),2)
            self.assertTrue(set(ids).isdisjoint(c['id'] for c in selected))
            self.assertEqual(sum(c['selection']=='backup' for c in after),4)
            self.assertTrue(set(ids).isdisjoint(c['id'] for c in after if c['selection']=='backup'))
            self.assertEqual(next(c for c in selected if c['id']==top[1]['id'])['selection'],'proposed')
            self.assertFalse(any(c['selection']=='validated' for c in after))
            other=services.create_campaign('Autre salon','2028',[{'name':top[0]['name'],'category':'Produits alimentaires'}],'test')
            self.assertEqual(services.campaign_results(other['id'])[1][0]['selection'],'shortlisted')
            services.set_selection_dismissed(campaign['id'],ids,dismissed=False)
            restored=services.campaign_results(campaign['id'])[1]
            self.assertEqual({c['id'] for c in restored if c['selection'] in ACTIVE},{c['id'] for c in top})
            ai.assert_not_called(); launch.assert_not_called()
        self.assertEqual(db.all_rows('companies',campaign['id']),before)
        self.assertEqual(db.all_rows('jobs'),jobs); self.assertEqual(db.all_rows('crm'),crm)
        self.assertTrue(all(updated['selection_decisions'][id]['dismissed'] for id in ids))

    def test_invalid_company_cannot_partially_change_campaign(self):
        before=db.get('campaigns',self.campaign['id']); audit=db.audit_rows()
        id=db.all_rows('companies',self.campaign['id'])[0]['id']
        with self.assertRaises(ValueError): services.set_selection_dismissed(self.campaign['id'],[id,'foreign'])
        self.assertEqual(db.get('campaigns',self.campaign['id']),before)
        self.assertEqual(db.audit_rows(),audit)

    def test_human_verified_and_legacy_proposals_also_respect_decisions(self):
        companies=db.all_rows('companies',self.campaign['id'])
        c=companies[0]
        for fact in c['facts'].values(): fact['status']='validated'
        criteria={**self.campaign['criteria'],'progressive_selection':False}
        self.assertEqual(next(r for r in rank(companies,criteria,[],crm_empty_confirmed=True) if r['id']==c['id'])['selection'],'validated')
        original=copy.deepcopy(companies)
        rows=rank(companies,criteria,[],crm_empty_confirmed=True,selection_decisions={c['id']:{'dismissed':True}})
        self.assertEqual(next(r for r in rows if r['id']==c['id'])['selection'],'none')
        self.assertEqual(companies,original)

    def test_all_export_formats_use_new_top_and_keep_dismissal_trace(self):
        from docx import Document
        from openpyxl import load_workbook
        campaign,before=services.campaign_results(self.campaign['id'])
        removed=[c for c in before if c['selection'] in ACTIVE][:2]
        services.set_selection_dismissed(campaign['id'],[c['id'] for c in removed])
        campaign,rows=services.campaign_results(campaign['id'])
        package=exports.prepare_package(campaign,rows,[],'Présélection automatique')
        expected={c['name'] for c in rows if c['selection'] in ACTIVE}
        self.assertEqual(len(expected),3); self.assertFalse(package['errors'])
        doc=Document(io.BytesIO(package['docx']))
        text='\n'.join(p.text for p in doc.paragraphs)
        for c in removed: self.assertNotIn(c['name'],text); self.assertNotIn(c['name'],package['csv'].decode('utf-8-sig'))
        for name in expected: self.assertIn(name,text)
        wb=load_workbook(io.BytesIO(package['xlsx']))
        def names(sheet): return {row[0] for row in wb[sheet].iter_rows(min_row=2,values_only=True)}
        self.assertEqual(names('Écartées de la sélection'),{c['name'] for c in removed})
        self.assertEqual(names('Propositions IA')|names('Présélection à documenter'),expected)
        self.assertTrue(expected.isdisjoint(names('Réserve')))
        with zipfile.ZipFile(io.BytesIO(package['zip'])) as z:
            result=json.loads(z.read('Resultats_et_sources.json'))
            self.assertEqual({c['name'] for c in result['companies']},expected)
            self.assertEqual(len([n for n in z.namelist() if n.startswith('fiches/') and n.endswith('.docx')]),3)
            self.assertEqual(z.read('Fiches_Memos.docx')[:2],b'PK')
            self.assertTrue(z.read('Fiches_Memos.pdf').startswith(b'%PDF'))

    def test_export_signature_tracks_selection_facts_crm_and_rules(self):
        campaign,rows=services.campaign_results(self.campaign['id'])
        signature=exports.export_signature(campaign,rows,[],'Présélection automatique')
        self.assertEqual(signature,exports.export_signature(copy.deepcopy(campaign),copy.deepcopy(rows),[],'Présélection automatique'))
        changed=copy.deepcopy(rows); changed[0]['memo']['pitch']='Nouveau pitch'
        self.assertNotEqual(signature,exports.export_signature(campaign,changed,[],'Présélection automatique'))
        self.assertNotEqual(signature,exports.export_signature(campaign,rows,[{'name':'Client'}],'Présélection automatique'))
        services.set_selection_dismissed(campaign['id'],[rows[0]['id']])
        updated,after=services.campaign_results(campaign['id'])
        self.assertNotEqual(signature,exports.export_signature(updated,after,[],'Présélection automatique'))


if __name__=='__main__': unittest.main(verbosity=2)
