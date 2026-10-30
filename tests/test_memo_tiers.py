import copy
import io
import unittest
from docx import Document
from pypdf import PdfReader

from impact import exports
from impact.config import DEFAULT_CRITERIA
from impact.domain import assess
from impact.signals import evidence_review,partition


def company(name,tier,score):
    return {'name':name,'facts':{},'memo':{},'contacts':[],'events':[],'triggers':[],
            'selection':'shortlisted','assessment':{'tier':tier,'score':score,'blockers':[]},
            'validated_assessment':{'tier':tier,'score':score,'blockers':[]}}


class MemoTierTests(unittest.TestCase):
    def test_word_pdf_and_export_selection_share_tier_order_and_visible_group_headings(self):
        companies=[company('Secondaire',3,45),company('Prioritaire',1,90),company('Intermédiaire',2,70),company('Prioritaire bis',1,85)]
        original=copy.deepcopy(companies)
        campaign={'name':'Salon de test','edition':'2027','criteria':DEFAULT_CRITERIA}
        selected=exports.selected_for_export(companies,'Présélection automatique')
        self.assertEqual([c['name'] for c in selected],['Prioritaire','Prioritaire bis','Intermédiaire','Secondaire'])
        doc=Document(io.BytesIO(exports.document_bytes(campaign,companies)))
        tier_headings=[p.text for p in doc.paragraphs if p.style.name=='Heading 1']
        self.assertEqual(tier_headings,[exports.tier_title(t) for t in [1,2,3]])
        titles=[p.text for p in doc.paragraphs if p.style.name=='Title']
        self.assertEqual(titles,[c['name'] for c in selected])
        text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(exports.pdf_bytes(campaign,companies))).pages)
        self.assertLess(text.index(exports.tier_title(1)),text.index(exports.tier_title(2)))
        self.assertLess(text.index(exports.tier_title(2)),text.index(exports.tier_title(3)))
        self.assertEqual(companies,original)

    def test_documented_mode_does_not_include_incomplete_shortlist(self):
        c=company('À compléter',3,45)
        self.assertEqual(exports.selected_for_export([c],'Critères documentés'),[])


class TriggerEvidenceTests(unittest.TestCase):
    def trigger(self,code,quote,status='sourced'):
        return {'code':code,'quote':quote,'source':'https://company.example/about','status':status,'observation':'Proposition du modèle.'}

    def test_explicit_examples_cover_all_eight_triggers(self):
        examples={'T1':'Un portefeuille diversifié de 150 produits et de nombreuses gammes.',
                  'T2':'Nos références de niche ont une faible rotation et une faible contribution.',
                  'T3':'Nos variantes de recettes et formats sont produites en petites séries.',
                  'T4':'Nos produits sont destinés à la restauration et à la grande distribution.',
                  'T5':'La croissance est faible et la pression sur les coûts augmente en 2025.',
                  'T6':'Après rationalisation des références, de nouvelles variantes apparaissent à nouveau.',
                  'T7':'Des gammes bio, sans gluten et certifiées équitables.',
                  'T8':'Nos best sellers se trouvent dans un catalogue très large et un portefeuille diversifié.'}
        for code,quote in examples.items():
            with self.subTest(code=code): self.assertTrue(evidence_review(self.trigger(code,quote))[0])

    def test_generic_catalogue_does_not_prove_low_rotation_and_premium_is_not_t7(self):
        self.assertFalse(evidence_review(self.trigger('T2','Nous développons des sauces, dips et produits végétaux sans additifs.'))[0])
        self.assertFalse(evidence_review(self.trigger('T7','Une gamme de caviars d’exception aux saveurs et qualités uniques.'))[0])
        self.assertFalse(evidence_review(self.trigger('T6','Nous innovons régulièrement et développons de nouvelles gammes.'))[0])
        self.assertFalse(evidence_review(self.trigger('T3','Notre expertise au service de vos créations d’exception depuis 1966.'))[0])
        self.assertFalse(evidence_review(self.trigger('T5','Notre croissance est forte et les coûts en hausse en 2025.'))[0])

    def test_real_catalogue_list_can_demonstrate_t1_without_a_prose_description(self):
        trigger=self.trigger('T1','Arla Pro Mozzarella râpée 2kg Best Seller\nArla Pro Cheddar tranches 1kg\nArla Pro Mozzarella tranches 1kg')
        self.assertTrue(evidence_review(trigger)[0])

    def test_unproven_triggers_remain_auditable_without_inflating_score(self):
        c=company('Test',3,0)
        c['triggers']=[self.trigger('T1','Un portefeuille diversifié de 150 produits et de nombreuses gammes.'),
                       self.trigger('T2','Notre expertise pour vos créations boulangères et pâtissières.'),
                       self.trigger('T7','Notre gamme premium aux saveurs d’exception.')]
        original=copy.deepcopy(c)
        a=assess(c,DEFAULT_CRITERIA,[],crm_empty_confirmed=True)
        self.assertEqual(a['components']['triggers'],0.25)
        self.assertEqual([t['code'] for t in a['trigger_review']],['T2','T7'])
        sections=dict(exports.memo_sections(c))
        self.assertTrue(any('T1' in s for s in sections['Signaux SENSO']))
        self.assertFalse(any('T2' in s for s in sections['Signaux SENSO']))
        self.assertTrue(any('T2' in s for s in sections['Signaux proposés à vérifier']))
        self.assertEqual(c,original)

    def test_human_validation_is_kept_and_confirmed_exports_skip_unreviewed_signals(self):
        c=company('Test',3,0)
        c['triggers']=[self.trigger('T2','Vérifié par la personne responsable.','validated'),
                       self.trigger('T7','Des gammes bio et certifiées équitables.')]
        self.assertEqual([t['code'] for t in partition(c,True)[0]],['T2'])
