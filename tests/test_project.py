import copy
import io
import json
import os
import sqlite3
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from impact import db,ai,services,exports
from impact.config import DEFAULT_CRITERIA,ROOT
from impact.domain import assess,rank,value,number,exclusion,contact_rows,validate_criteria,research_key
from impact.importers import parse_file,normalize_records,table

YEAR=datetime.now().year

def qualifying(name='Exemple Alimentaire'):
    c={'id':db.uid(),'campaign_id':'test','name':name,'facts':{},'contacts':[], 'triggers':[], 'events':[], 'memo':{}, 'warnings':[], 'source_rows':[],'match_decisions':{}}
    fields={'identity':True,'sector_fit':True,'business_model':'industriel','international':True,
            'revenue_eur':12000000.0,'finance_year':YEAR-1,'finance_scope':'entite','city':'Lyon','stand':'A12','hall':'2'}
    for f,v in fields.items(): c['facts'][f]={'value':v,'status':'validated','source':'https://example.com','quote':'exemple pour test','year':YEAR-1,'scope':'entite'}
    return c

CRM=[{'id':'crm1','name':'Ancien Client','website':'https://ancien-client.example','siren':'','status':'client'}]

class ProjectTests(unittest.TestCase):
    def setUp(self):
        temp=tempfile.TemporaryDirectory(prefix='impact_project_test_')
        self.addCleanup(temp.cleanup)
        data=patch.object(db,'DATA',Path(temp.name))
        data.start(); self.addCleanup(data.stop)
        db.init()
        self.criteria=copy.deepcopy(DEFAULT_CRITERIA)
        self.criteria['progressive_selection']=False  # Contrat du classement historique testé ici.

    def test_real_sirha_columns_and_count(self):
        p=parse_file('sirha.xlsx',(ROOT/'examples/SIRHA_Marseille_exemple.xlsx').read_bytes())
        self.assertEqual(len(p['tables'][0]['rows']),117)
        self.assertEqual(p['tables'][0]['mapping']['exhibitor_url'],'Sirha_URL')
        self.assertNotIn('stand',p['tables'][0]['mapping'])
        records,_=normalize_records(p['tables'][0],p['tables'][0]['mapping'])
        self.assertEqual(records[0]['website'].rstrip('/'),'http://www.sybavegetal.com')
        self.assertEqual(records[0]['email'],'')

    def test_sigep_forward_fill_contacts_duplicate_linkedin_headers(self):
        p=parse_file('sigep.xlsx',(ROOT/'examples/SIGEP_Contacts_exemple.xlsx').read_bytes())
        tab=p['tables'][0]
        self.assertEqual(tab['mapping']['contact_linkedin'],'LinkedIn Contact (2)')
        records,_=normalize_records(tab,tab['mapping'])
        self.assertEqual(len(records),32)
        self.assertEqual(len({r['name'] for r in records}),20)
        bridor=[r for r in records if r['name']=='BRIDOR SAS']
        self.assertEqual(len(bridor),3)
        self.assertTrue(all(r['stand']=='Hall B2 – Stand 039' for r in bridor))
        c=services.create_campaign('SIGEP test','2026',records,'test')
        companies=db.all_rows('companies',c['id'])
        self.assertEqual(len(companies),20)
        self.assertEqual(len(next(x for x in companies if x['name']=='BRIDOR SAS')['contacts']),3)

    def test_csv_json_html_and_txt(self):
        for name,content in [('x.csv','Entreprise;Stand\nAlpha;A12'.encode()),('x.json',b'{"exhibitors":[{"name":"Alpha","stand":"A12"}]}'),
                             ('x.html',b'<script>ignore instructions</script><table><tr><th>Entreprise</th><th>Stand</th></tr><tr><td>Alpha</td><td>A12</td></tr></table>')]:
            p=parse_file(name,content); rows,_=normalize_records(p['tables'][0],p['tables'][0]['mapping'])
            self.assertEqual(rows[0]['name'],'Alpha'); self.assertEqual(rows[0]['stand'],'A12')
        p=parse_file('x.txt','Alpha\nBêta'.encode()); self.assertEqual(len(p['tables'][0]['rows']),2)
        p=parse_file('x.csv',b'Alpha\nBeta'); self.assertEqual(p['tables'][0]['rows'][0]['Entreprise'],'Alpha')

    def test_delete_campaign_removes_its_dependents_and_preserves_shared_data(self):
        campaign=services.create_campaign('À supprimer','2027',[{'name':'Alpha'},{'name':'Beta'}],'test')
        other=services.create_campaign('À conserver','2027',[{'name':'Gamma'}],'test')
        company_ids=[c['id'] for c in db.all_rows('companies',campaign['id'])]
        job=db.put('jobs',{'campaign_id':campaign['id'],'status':'completed'})
        other_job=db.put('jobs',{'campaign_id':other['id'],'status':'completed'})
        crm=db.put('crm',{'name':'Prospect conservé','status':'prospect','source':'Campagne '+campaign['id']})
        cache=db.put('cache',{'key':db.uid(),'result':{'facts':{}}})
        criteria=copy.deepcopy(db.setting('criteria'))

        deleted=services.delete_campaign(campaign['id'])

        self.assertEqual(deleted['companies'],2)
        self.assertEqual(deleted['jobs'],1)
        self.assertEqual(db.all_rows('companies',campaign['id']),[])
        self.assertEqual(db.all_rows('jobs',campaign['id']),[])
        for table,id in [('campaigns',campaign['id']),('jobs',job['id']),*[('companies',id) for id in company_ids]]:
            with self.assertRaises(ValueError): db.get(table,id)
        self.assertEqual(db.get('campaigns',other['id'])['name'],'À conserver')
        self.assertEqual(len(db.all_rows('companies',other['id'])),1)
        self.assertEqual(db.get('jobs',other_job['id'])['status'],'completed')
        self.assertEqual(db.get('crm',crm['id'])['name'],'Prospect conservé')
        self.assertEqual(db.get('cache',cache['key'])['result'],{'facts':{}})
        self.assertEqual(db.setting('criteria'),criteria)
        self.assertTrue(any(r['action']=='campaign_delete' and json.loads(r['payload'])['campaign_id']==campaign['id'] for r in db.audit_rows()))
        with self.assertRaisesRegex(ValueError,'n’existe plus'): services.delete_campaign(campaign['id'])

    def test_delete_campaign_blocks_queued_or_running_jobs_even_after_cancel_request(self):
        campaign=services.create_campaign('En cours','2027',[{'name':'Alpha'}],'test')
        job=db.put('jobs',{'campaign_id':campaign['id'],'status':'queued','cancel':True})
        for status in ['queued','running']:
            with self.subTest(status=status):
                job['status']=status; db.put('jobs',job)
                with self.assertRaisesRegex(ValueError,'encore en cours'): services.delete_campaign(campaign['id'])
                self.assertEqual(db.get('campaigns',campaign['id'])['name'],'En cours')
                self.assertEqual(len(db.all_rows('companies',campaign['id'])),1)
                self.assertEqual(db.get('jobs',job['id'])['status'],status)
        job['status']='cancelled'; db.put('jobs',job)
        services.delete_campaign(campaign['id'])

    def test_delete_campaign_rolls_back_on_database_failure(self):
        campaign=services.create_campaign('Transaction','2027',[{'name':'Alpha'}],'test')
        job=db.put('jobs',{'campaign_id':campaign['id'],'status':'interrupted'})
        with db.connection() as conn:
            conn.execute("CREATE TRIGGER reject_campaign_delete BEFORE DELETE ON campaigns BEGIN SELECT RAISE(ABORT,'delete rejected'); END")
        try:
            with self.assertRaises(sqlite3.IntegrityError): services.delete_campaign(campaign['id'])
            self.assertEqual(db.get('campaigns',campaign['id'])['name'],'Transaction')
            self.assertEqual(len(db.all_rows('companies',campaign['id'])),1)
            self.assertEqual(db.get('jobs',job['id'])['status'],'interrupted')
            self.assertFalse(any(r['action']=='campaign_delete' and json.loads(r['payload'])['campaign_id']==campaign['id'] for r in db.audit_rows()))
        finally:
            with db.connection() as conn: conn.execute('DROP TRIGGER reject_campaign_delete')

    def test_job_research_persists_and_preserves_validated_fields(self):
        campaign=services.create_campaign('Job test','2027',[{'name':'Job Alpha','website':'https://alpha.example'}],'test')
        c=db.all_rows('companies',campaign['id'])[0]
        c['facts']['city']={'value':'Paris','status':'validated','source':'Saisie utilisateur','quote':'Paris'}
        db.put('companies',c)
        job={'id':db.uid(),'campaign_id':campaign['id'],'created':db.now(),'status':'queued','done':0,'total':1,
             'message':'','errors':[],'cancel':False,'criteria_snapshot':self.criteria,'company_ids':[c['id']],'refresh':True,'model':'mock'}
        db.put('jobs',job)
        result={'facts':{'city':{'value':'Lyon','status':'sourced','source':'https://alpha.example','quote':'Implantation située à Lyon'}},
                'contacts':[],'triggers':[],'events':[],'memo':{},'warnings':[],'identity_match':True,'sources':[],'usage':{}}
        with patch('impact.ai.research',return_value=result): services.run_job(job['id'])
        self.assertEqual(db.get('jobs',job['id'])['status'],'completed')
        self.assertEqual(value(db.get('companies',c['id']),'city'),'Paris')

    def test_job_cancel_is_persisted_and_no_call_is_made(self):
        campaign=services.create_campaign('Cancel test','2027',[{'name':'Cancel Alpha'}],'test')
        c=db.all_rows('companies',campaign['id'])[0]
        job={'id':db.uid(),'campaign_id':campaign['id'],'created':db.now(),'status':'queued','done':0,'total':1,
             'message':'','errors':[],'cancel':True,'criteria_snapshot':self.criteria,'company_ids':[c['id']],'refresh':True,'model':'mock'}
        db.put('jobs',job)
        with patch('impact.ai.research') as mocked: services.run_job(job['id'])
        self.assertFalse(mocked.called); self.assertEqual(db.get('jobs',job['id'])['status'],'cancelled')

    def test_fatal_api_error_does_not_run_remaining_companies(self):
        campaign=services.create_campaign('Error test','2027',[{'name':'Error Alpha'},{'name':'Error Beta'}],'test')
        ids=[x['id'] for x in db.all_rows('companies',campaign['id'])]
        job={'id':db.uid(),'campaign_id':campaign['id'],'created':db.now(),'status':'queued','done':0,'total':2,
             'message':'','errors':[],'cancel':False,'criteria_snapshot':self.criteria,'company_ids':ids,'refresh':True,'model':'mock'}
        db.put('jobs',job)
        with patch('impact.ai.research',side_effect=ai.AIError('Clé API refusée.')) as mocked: services.run_job(job['id'])
        self.assertEqual(mocked.call_count,1); self.assertEqual(db.get('jobs',job['id'])['status'],'failed')

    def test_missing_name_does_not_forward_fill_without_contact(self):
        tab=table('x',[['Entreprise','Stand'],['Alpha','1'],['','2'],['Beta','3']])
        rows,warnings=normalize_records(tab,tab['mapping'])
        self.assertEqual(len(rows),2); self.assertEqual(len(warnings),1)

    def test_numeric_locale_units_and_zero(self):
        for v in ['12 500 000','12,5 M','12.5M','12,500,000','12.500.000','12 500 000 EUR']:
            self.assertEqual(number(v),12500000)
        self.assertEqual(number('0'),0); self.assertIsNone(number('10 à 15 M'))

    def test_financial_threshold_is_strict_and_profit_alternative(self):
        c=qualifying(); c['facts']['revenue_eur']['value']=10000000
        self.assertFalse(assess(c,self.criteria,CRM,True)['eligible'])
        c['facts']['margin_pct']={'value':10,'status':'validated','year':YEAR-1,'scope':'entite'}
        self.assertTrue(assess(c,self.criteria,CRM,True)['eligible'])

    def test_unknown_stale_group_and_future_financials(self):
        c=qualifying()
        for year,scope in [(None,None),(YEAR-10,'entite'),(YEAR+1,'entite'),(YEAR-1,'groupe')]:
            c['facts']['revenue_eur'].update(year=year,scope=scope)
            if year is None:
                c['facts'].pop('finance_year',None); c['facts'].pop('finance_scope',None)
            self.assertFalse(assess(c,self.criteria,CRM,True)['eligible'])
        self.criteria['allow_group_finance']=True
        c['facts']['entity_type']={'value':'marque','status':'validated'}
        self.assertFalse(assess(c,self.criteria,CRM,True)['eligible'])
        c['facts']['entity_type']['value']='groupe'
        self.assertTrue(assess(c,self.criteria,CRM,True)['eligible'])

    def test_empty_crm_prevents_final_qualification(self):
        c=qualifying(); self.assertFalse(assess(c,self.criteria,[],True)['eligible'])
        ranked=rank([c],self.criteria,[]); self.assertEqual(ranked[0]['selection'],'proposed')

    def test_imported_or_sourced_facts_do_not_become_validated(self):
        c=qualifying()
        for f in c['facts'].values(): f['status']='sourced'
        self.assertTrue(assess(c,self.criteria,CRM)['eligible'])
        self.assertFalse(assess(c,self.criteria,CRM,True)['eligible'])
        self.assertEqual(rank([c],self.criteria,CRM)[0]['selection'],'proposed')

    def test_exact_crm_domain_siren_name_and_fuzzy(self):
        c=qualifying('Ancien Client SAS')
        self.assertEqual(exclusion(c,CRM)['state'],'excluded')
        c['name']='Nom différent'; c['facts']['website']={'value':'https://www.ancien-client.example/contact'}
        self.assertEqual(exclusion(c,CRM)['state'],'excluded')
        c['facts'].pop('website'); c['facts']['siren']={'value':'12345678900001'}
        self.assertEqual(exclusion(c,[{**CRM[0],'siren':'123456789'}])['state'],'excluded')
        c['facts'].pop('siren'); c['name']='Ancien Clients'
        self.assertEqual(exclusion(c,CRM)['state'],'review')
        c['match_decisions']={'crm1':{'action':'include','reason':'entité distincte'}}
        self.assertEqual(exclusion(c,CRM)['state'],'clear')
        self.assertEqual(exclusion(qualifying('BINDI S.p.A.'),[{'id':'x','name':'Bindi','status':'prospect'}])['state'],'excluded')

    def test_aggregate_preserves_existing_and_qualified_contacts(self):
        old=[{'id':'old','name':'Client historique','status':'client','email':'info@client.example',
              'contacts':[{'name':'Contact ancien','email':'person@client.example'}]}]
        c=qualifying(); c['contacts']=[{'name':'Contact nouveau','email':'person@new.example','status':'validated'}]
        rows=exports.aggregate_rows(old,rank([c],self.criteria,CRM))
        self.assertEqual(len(rows),2); self.assertEqual(rows[0]['Email entreprise'],'info@client.example'); self.assertEqual(rows[1]['Email contact'],'person@new.example')

    def test_top_twenty_stable_and_never_select_excluded(self):
        companies=[qualifying('Entreprise '+str(i)) for i in range(25)]
        companies.append(qualifying('Ancien Client'))
        ranked=rank(companies,self.criteria,CRM)
        self.assertEqual(sum(c['selection']=='validated' for c in ranked),20)
        self.assertEqual(next(c for c in ranked if c['name']=='Ancien Client')['selection'],'none')

    def test_sector_change_invalidates_previous_interpretation(self):
        c=qualifying(); c['facts']['sector_fit']['criteria_sectors']=self.criteria['sectors']
        self.criteria['sectors']=['cosmétique']
        self.assertFalse(assess(c,self.criteria,CRM,True)['eligible'])

    def test_criteria_bad_weights_rejected(self):
        self.criteria['weights']['sector']=50
        with self.assertRaises(ValueError): validate_criteria(self.criteria)

    def test_stand_gating_and_tiers(self):
        c=qualifying(); self.assertEqual(assess(c,self.criteria,CRM,True)['tier'],1)
        c['facts'].pop('stand'); self.criteria['require_stand']=True
        self.assertFalse(assess(c,self.criteria,CRM,True)['eligible'])

    def test_ai_rejects_missing_sources_guessed_emails_wrong_edition(self):
        source='https://example.com/contact'; campaign={'name':'Salon','edition':'2027'}
        result={'identity_match':True,'facts':[
            {'field':'email','value':'guessed@example.com','source_url':source,'quote':'Notre équipe commerciale vous accompagne'},
            {'field':'revenue_eur','value':'15000000','source_url':'https://not-consulted.example','quote':'15 millions de chiffre affaires'},
            {'field':'stand','value':'A12','edition':'2026','source_url':source,'quote':'Stand A12 pour cette édition de salon'}],
            'contacts':[{'name':'Alice Test','role':'Export','email':'alice@example.com','phone':'','linkedin':'','source_url':source,'quote':'Alice Test dirige notre équipe export'}]}
        filtered=ai.filter_result(result,{source},campaign,self.criteria)
        self.assertFalse(filtered['facts']); self.assertEqual(filtered['contacts'][0]['email'],'')

    def test_ai_contradiction_not_arbitrarily_resolved(self):
        source='https://example.com'
        result={'identity_match':True,'facts':[{'field':'city','value':v,'source_url':source,'quote':'Implantation à la ville concernée'} for v in ['Paris','Lyon','Paris']]}
        self.assertNotIn('city',ai.filter_result(result,{source},{'edition':'2026'},self.criteria)['facts'])

    def test_ai_unknown_identity_blocks_all_facts(self):
        filtered=ai.filter_result({'identity_match':False,'identity_explanation':'deux sociétés homonymes','facts':[{}]},set(),{'edition':'2026'},self.criteria)
        self.assertFalse(filtered['facts']); self.assertFalse(filtered['identity_match'])

    def test_api_contract_two_calls_mocked(self):
        source='https://example.com'
        search={'status':'completed','output':[{'type':'web_search_call','action':{'sources':[{'url':source}]}},
                     {'type':'message','content':[{'type':'output_text','text':'Recherche source','annotations':[]}]}]}
        result={'identity_match':True,'identity_explanation':'Entreprise confirmée','facts':[],'contacts':[],'triggers':[],'events':[],'pitch':'','pain_points_hypotheses':[],'questions':[],'key_message':'','missing_information':[]}
        extraction={'status':'completed','output':[{'type':'message','content':[{'type':'output_text','text':json.dumps(result)}]}]}
        with patch.dict(os.environ,{'OPENAI_API_KEY':'fake-test-key','AI_PROVIDER':'openai'}),patch('impact.ai.call',side_effect=[search,extraction]) as mocked:
            out=ai.research(qualifying(),{'name':'Salon','edition':'2027'},self.criteria)
        self.assertEqual(mocked.call_count,2)
        self.assertEqual(mocked.call_args_list[0].args[0]['tools'],[{'type':'web_search'}])
        self.assertTrue(mocked.call_args_list[1].args[0]['text']['format']['strict'])
        self.assertEqual(out['sources'],[source])

    def test_cache_includes_edition_criteria_and_knowledge(self):
        c=qualifying(); campaign={'name':'Salon','edition':'2026'}
        key=research_key(c,campaign,self.criteria,'model')
        campaign['edition']='2027'; self.assertNotEqual(key,research_key(c,campaign,self.criteria,'model'))

    def test_export_no_formula_injection_and_no_guessed_contacts(self):
        from openpyxl import load_workbook
        c=qualifying('=HYPERLINK("http://bad")'); ranked=rank([c],self.criteria,CRM)
        campaign={'name':'Salon','edition':'2027','criteria':self.criteria}
        output=exports.workbook_bytes(campaign,ranked,CRM)
        wb=load_workbook(io.BytesIO(output))
        self.assertNotEqual(wb['Cibles validées']['A2'].data_type,'f')
        self.assertTrue(wb['Cibles validées']['A2'].value.startswith("'="))
        rows=contact_rows(ranked,True); self.assertEqual(rows[0]['Email'],'')
        self.assertIsInstance(wb['Cibles validées']['I2'].value,(float,int))

    def test_docx_pdf_bundle_have_stand_and_sources(self):
        from docx import Document
        from pypdf import PdfReader
        c=qualifying(); ranked=rank([c],self.criteria,CRM)
        campaign={'name':'Salon','edition':'2027','criteria':self.criteria}
        doc=Document(io.BytesIO(exports.document_bytes(campaign,ranked,True)))
        self.assertIn('STAND A12','\n'.join(p.text for p in doc.paragraphs))
        pdf=PdfReader(io.BytesIO(exports.pdf_bytes(campaign,ranked,True)))
        self.assertIn('STAND A12','\n'.join(p.extract_text() for p in pdf.pages))
        with zipfile.ZipFile(io.BytesIO(exports.bundle_bytes(campaign,ranked,CRM,True))) as z:
            self.assertIn('Fiches_Memos.docx',z.namelist()); self.assertEqual(len([x for x in z.namelist() if x.startswith('fiches/')]),2)

    def test_no_key_never_calls_network(self):
        with patch.dict(os.environ,{'OPENAI_API_KEY':''}),patch('impact.ai.request.urlopen') as mocked:
            with self.assertRaises(ai.AIError): ai.call({'input':'x'})
            self.assertFalse(mocked.called)

    def test_save_review_preserves_year_scope_and_zero(self):
        campaign=services.create_campaign('Review test','2027',[{'name':'Test Review','revenue_eur':'0'}],'source')
        c=db.all_rows('companies',campaign['id'])[0]
        services.save_review(c['id'],{'margin_pct':'0','finance_year':str(YEAR-1),'finance_scope':'entite'},[],[],[],{'margin_pct','finance_year','finance_scope'})
        c=db.get('companies',c['id']); self.assertEqual(value(c,'margin_pct'),0); self.assertEqual(c['facts']['margin_pct']['year'],YEAR-1)

if __name__=='__main__': unittest.main(verbosity=2)
